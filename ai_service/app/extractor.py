import io
import re
from typing import Dict, List, Tuple, Any
# pyrefly: ignore [missing-import]
from pypdf import PdfReader
from docx import Document

# Extensive Skills Taxonomy
SKILL_TAXONOMY: Dict[str, List[str]] = {
    "Programming Languages": [
        "python", "javascript", "typescript", "java", "c++", "c#", "c", "go", "golang",
        "ruby", "php", "swift", "kotlin", "rust", "scala", "r", "dart", "shell", "bash",
        "powershell", "perl", "matlab", "sql", "html", "css", "sass", "scss"
    ],
    "Frontend Frameworks & Web": [
        "react", "react.js", "reactjs", "next.js", "nextjs", "vue", "vue.js", "vuejs",
        "angular", "angularjs", "svelte", "redux", "zustand", "tailwind", "tailwindcss",
        "bootstrap", "material-ui", "mui", "chakra ui", "jquery", "webpack", "vite",
        "html5", "css3", "responsive design", "web design", "rest api"
    ],
    "Backend & APIs": [
        "node.js", "nodejs", "express", "express.js", "fastapi", "django", "flask",
        "spring", "spring boot", "asp.net", ".net core", "graphql", "restful apis",
        "microservices", "nest.js", "nestjs", "socket.io", "grpc", "celery"
    ],
    "Databases & Caching": [
        "mongodb", "postgresql", "postgres", "mysql", "sqlite", "redis", "cassandra",
        "elasticsearch", "dynamodb", "mariadb", "oracle", "prisma", "mongoose", "sequelize"
    ],
    "Cloud & DevOps": [
        "docker", "kubernetes", "k8s", "aws", "amazon web services", "azure", "google cloud", "gcp",
        "ci/cd", "github actions", "gitlab ci", "jenkins", "terraform", "ansible",
        "nginx", "apache", "linux", "git", "github", "bitbucket"
    ],
    "AI, Machine Learning & Data Science": [
        "machine learning", "deep learning", "nlp", "natural language processing",
        "computer vision", "tensorflow", "pytorch", "keras", "scikit-learn", "sklearn",
        "pandas", "numpy", "matplotlib", "seaborn", "huggingface", "transformers",
        "opencv", "llm", "genai", "prompt engineering", "data analysis", "data mining"
    ],
    "Tools & Practices": [
        "agile", "scrum", "jira", "postman", "jest", "pytest", "unit testing",
        "cypress", "selenium", "figma", "vs code", "pycharm", "docker compose"
    ]
}

# Flattened list for fast normalized lookup
ALL_SKILLS: Dict[str, str] = {}
for category, skills in SKILL_TAXONOMY.items():
    for s in skills:
        ALL_SKILLS[s.lower()] = s

# Aliases mapping (alias -> canonical skill)
SKILL_ALIASES: Dict[str, str] = {
    "reactjs": "react",
    "react.js": "react",
    "nodejs": "node.js",
    "node": "node.js",
    "nextjs": "next.js",
    "expressjs": "express",
    "express.js": "express",
    "golang": "go",
    "postgres": "postgresql",
    "sklearn": "scikit-learn",
    "aws": "amazon web services",
    "k8s": "kubernetes",
    "tail-wind": "tailwind",
    "spring-boot": "spring boot",
    "ml": "machine learning",
    "dl": "deep learning"
}

DEGREE_PATTERNS = [
    r"\b(ph\.?d|doctor of philosophy)\b",
    r"\b(m\.?tech|master of technology)\b",
    r"\b(m\.?s|master of science)\b",
    r"\b(m\.?c\.?a|master of computer applications)\b",
    r"\b(m\.?b\.?a|master of business administration)\b",
    r"\b(b\.?tech|bachelor of technology)\b",
    r"\b(b\.?e|bachelor of engineering)\b",
    r"\b(b\.?s|bachelor of science)\b",
    r"\b(b\.?c\.?a|bachelor of computer applications)\b",
    r"\b(b\.?sc|bachelor of science)\b",
    r"\b(diploma in [a-zA-Z\s]+|polytechnic)\b",
    r"\b(high school|secondary education|12th grade)\b"
]

def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract full raw text from PDF bytes."""
    text = ""
    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        for page in reader.pages:
            extracted = page.extract_text()
            if extracted:
                text += extracted + "\n"
    except Exception as e:
        text = f"Error reading PDF: {str(e)}"
    return text.strip()

def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract text from DOCX file bytes."""
    text = ""
    try:
        doc = Document(io.BytesIO(file_bytes))
        for para in doc.paragraphs:
            if para.text.strip():
                text += para.text + "\n"
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    text += " | ".join(row_text) + "\n"
    except Exception as e:
        text = f"Error reading DOCX: {str(e)}"
    return text.strip()

def extract_email(text: str) -> str:
    """Extract primary candidate email."""
    pattern = r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+"
    matches = re.findall(pattern, text)
    if matches:
        return matches[0].lower().strip()
    return ""

def extract_phone(text: str) -> str:
    """Extract primary contact phone number."""
    pattern = r"(?:\+?\d{1,3}[-.\s]?)?(?:\d{5}[-.\s]?\d{5}|\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}|\b\d{10}\b)"
    matches = re.findall(pattern, text)
    if matches:
        return matches[0].strip()
    return ""

def extract_candidate_name(text: str) -> str:
    """Extract candidate name from the top header lines of the resume."""
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    for line in lines[:8]:
        # Filter out email, phone, links, or common header words
        if "@" in line or "http" in line or "resume" in line.lower() or "curriculum" in line.lower() or "summary" in line.lower():
            continue
        # If line contains title separator e.g. "Praveen Kumar - Full Stack Developer"
        segment = re.split(r"[-–|:]", line)[0].strip()
        clean_name = re.sub(r"[^a-zA-Z\s]", "", segment).strip()
        words = clean_name.split()
        if 1 <= len(words) <= 4 and all(w[0].isupper() for w in words if w):
            return clean_name
    return "Candidate"

def extract_education(text: str) -> str:
    """Detect highest education qualification mentioned in text."""
    lower_text = text.lower()
    for pattern in DEGREE_PATTERNS:
        match = re.search(pattern, lower_text)
        if match:
            matched_str = match.group(0).upper()
            if "PH" in matched_str:
                return "Ph.D. / Doctorate"
            elif "M.TECH" in matched_str or "MASTER OF TECH" in matched_str:
                return "M.Tech / Master of Technology"
            elif "M.S" in matched_str or "MASTER OF SCIENCE" in matched_str:
                return "M.S. / Master of Science"
            elif "MCA" in matched_str or "M.C.A" in matched_str:
                return "MCA / Master of Computer Applications"
            elif "MBA" in matched_str:
                return "MBA / Master of Business Administration"
            elif "B.TECH" in matched_str or "BACHELOR OF TECH" in matched_str:
                return "B.Tech / Bachelor of Technology"
            elif "B.E" in matched_str or "BACHELOR OF ENG" in matched_str:
                return "B.E. / Bachelor of Engineering"
            elif "B.S" in matched_str or "B.SC" in matched_str:
                return "B.S. / B.Sc in Computer Science"
            elif "BCA" in matched_str or "B.C.A" in matched_str:
                return "BCA / Bachelor of Computer Applications"
            elif "DIPLOMA" in matched_str:
                return "Diploma / Polytechnic"
            return matched_str
    return "Bachelor's Degree in Computer Science / Engineering"

def extract_experience_years(text: str) -> float:
    """Extract candidate total years of experience using regex patterns."""
    lower_text = text.lower()
    
    # Check explicit mentions e.g. "3.5 years of experience", "4+ yrs experience"
    exp_patterns = [
        r"(\d+(?:\.\d+)?)\s*\+?\s*(?:years?|yrs?)(?:\s+of)?\s+experience",
        r"experience\s*:\s*(\d+(?:\.\d+)?)\s*\+?\s*(?:years?|yrs?)",
        r"total\s+experience\s*:\s*(\d+(?:\.\d+)?)\s*\+?\s*(?:years?|yrs?)"
    ]
    for pattern in exp_patterns:
        match = re.search(pattern, lower_text)
        if match:
            try:
                return float(match.group(1))
            except ValueError:
                pass
                
    # Detect year ranges like 2021 - 2024 or 2020 - Present
    range_pattern = r"\b(20[0-2]\d)\s*(?:-|to|–)\s*(20[0-2]\d|present|current)\b"
    ranges = re.findall(range_pattern, lower_text)
    total_calculated_years = 0.0
    current_year = 2026
    
    for start_year, end_year in ranges:
        try:
            start = int(start_year)
            end = current_year if ("present" in end_year or "current" in end_year) else int(end_year)
            if end >= start and (end - start) <= 15:
                total_calculated_years += (end - start)
        except ValueError:
            pass
            
    if total_calculated_years > 0:
        return min(total_calculated_years, 20.0)
        
    return 1.0  # Default baseline for entry-level/fresher with projects

def extract_skills(text: str) -> Tuple[List[str], Dict[str, List[str]]]:
    """Extract recognized skills and return both flattened list and categorized map."""
    lower_text = text.lower()
    # Normalize punctuation around words
    clean_text = " " + re.sub(r"[,/|()\[\]•;:]", " ", lower_text) + " "
    
    found_skills_set = set()
    categorized_skills: Dict[str, List[str]] = {}
    
    for category, skill_list in SKILL_TAXONOMY.items():
        matched_in_cat = []
        for skill in skill_list:
            s_lower = skill.lower()
            # Boundary-safe regex search
            escaped = re.escape(s_lower)
            pattern = rf"(?:\b|\s){escaped}(?:\b|\s)"
            if re.search(pattern, clean_text):
                canonical = ALL_SKILLS.get(s_lower, skill)
                # Map alias if exists
                canonical = SKILL_ALIASES.get(canonical.lower(), canonical)
                found_skills_set.add(canonical)
                if canonical not in matched_in_cat:
                    matched_in_cat.append(canonical)
        if matched_in_cat:
            categorized_skills[category] = sorted(list(set(matched_in_cat)))
            
    # Sort and return unique list
    unique_skills = sorted(list(found_skills_set), key=lambda x: x.lower())
    return unique_skills, categorized_skills

def parse_resume_content(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    """Parse resume from raw bytes according to file type."""
    lower_name = filename.lower()
    if lower_name.endswith(".pdf"):
        raw_text = extract_text_from_pdf(file_bytes)
    elif lower_name.endswith(".docx"):
        raw_text = extract_text_from_docx(file_bytes)
    else:
        # Fallback to UTF-8 decoded text
        raw_text = file_bytes.decode("utf-8", errors="ignore")
        
    candidate_name = extract_candidate_name(raw_text)
    email = extract_email(raw_text)
    phone = extract_phone(raw_text)
    education = extract_education(raw_text)
    experience_years = extract_experience_years(raw_text)
    skills, categorized = extract_skills(raw_text)
    
    # Generate brief summary
    summary = f"{candidate_name} holds {education} with ~{experience_years:g} years of experience and proficiencies in {', '.join(skills[:5]) if skills else 'Software Development'}."
    
    return {
        "candidate_name": candidate_name,
        "email": email,
        "phone": phone,
        "education": education,
        "experience_years": experience_years,
        "skills": skills,
        "categorized_skills": categorized,
        "summary": summary,
        "raw_text": raw_text[:5000]  # Store first 5000 chars
    }
